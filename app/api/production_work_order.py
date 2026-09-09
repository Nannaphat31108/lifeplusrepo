import json
import re
from datetime import date
from io import BytesIO
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.session import get_db
from app.core.security import get_current_user, require_roles
from app.core.record_versions import snapshot_version, serialize_version
from app.models.entities import ProductionWorkOrder, RecordVersion

router = APIRouter(prefix="/api/production-work-orders", tags=["Production Work Order (PLANNING)"])

LOGO_PATH = Path(__file__).resolve().parents[1] / "static" / "logo.png"

# The source document's real 10-step workflow, pre-filled as a starting
# point for a new order (still fully editable: add/remove/retype any row)
# -- not invented, copied from the actual "ใบสั่งผลิต" sheet's own rows.
DEFAULT_WORKFLOW_STEPS = [
    {"task": "เช็คและขอซื้อสารสกัดและวัตถุดิบ", "responsible": "จนท.สต็อก"},
    {"task": "การแก้ไขหรือออกแบบบรรจุภัณฑ์", "responsible": "จนท.กราฟฟิก"},
    {"task": "การสั่งซื้อบรรจุภัณฑ์", "responsible": "จนท.จัดซื้อ(ทั่วไป)"},
    {"task": "การสั่งซื้อสารสกัด", "responsible": "จนท.จัดซื้อ(สารสกัด)"},
    {"task": "การตักสารและเตรียมเบิกวัตถุดิบ", "responsible": "หัวหน้าสต็อก"},
    {"task": "งานเทสสินค้าก่อนผลิต", "responsible": "จนท. RD"},
    {"task": "สรุปและแก้ปัญหากรณีการเทสไม่ผ่าน", "responsible": "จนท. QA"},
    {"task": "การผลิตสินค้า", "responsible": "จนท. Planning ฝ่ายผลิต"},
    {"task": "การโอนชำระสินค้าก่อนส่งสินค้า", "responsible": "จนท.ฝ่ายขาย"},
    {"task": "การจัดส่งสินค้า", "responsible": "จนท.จัดส่ง"},
]


def _add_logo_if_present(ws, anchor: str, width: int = 170, height: int = 73) -> None:
    if not LOGO_PATH.exists():
        return
    try:
        from openpyxl.drawing.image import Image as XLImage
        img = XLImage(str(LOGO_PATH))
        img.width, img.height = width, height
        ws.add_image(img, anchor)
    except Exception:
        pass


class OrderSave(BaseModel):
    order_no: str
    status: str = "DRAFT"
    data: dict
    linked_reference: str | None = None


def _validate_order_data(data: dict):
    if not str(data.get("product_name") or "").strip():
        raise HTTPException(400, "กรุณาใส่ชื่อผลิตภัณฑ์ก่อนบันทึก")
    groups = data.get("packing_groups") or []
    has_item = any(
        str(it.get("description") or "").strip()
        for g in groups for it in (g.get("items") or [])
    )
    if groups and not has_item:
        raise HTTPException(400, "กรุณาใส่รายการบรรจุภัณฑ์อย่างน้อย 1 รายการ ก่อนบันทึก")


def _serialize(x: ProductionWorkOrder) -> dict:
    return {
        "id": x.id,
        "order_no": x.order_no,
        "status": x.status,
        "data": json.loads(x.payload_json or "{}"),
        "linked_reference": x.linked_reference,
        "created_by_name": x.created_by_name,
        "created_at": x.created_at,
        "updated_at": x.updated_at,
    }


@router.get("/default-workflow")
def default_workflow(u=Depends(get_current_user)):
    return DEFAULT_WORKFLOW_STEPS


@router.post("")
def save_order(
    p: OrderSave,
    db: Session = Depends(get_db),
    u=Depends(get_current_user),
):
    if not p.data.get("date"):
        p.data["date"] = date.today().isoformat()
    _validate_order_data(p.data)

    x = ProductionWorkOrder(
        order_no=p.order_no,
        status=p.status,
        payload_json=json.dumps(p.data, ensure_ascii=False, default=str),
        created_by=u.id,
        created_by_name=u.full_name,
        linked_reference=(p.linked_reference or "").strip() or None,
    )
    db.add(x); db.commit(); db.refresh(x)
    return _serialize(x)


@router.get("")
def list_orders(db: Session = Depends(get_db), u=Depends(get_current_user)):
    """Department-shared listing -- same visibility model as PO/PR: any
    authenticated user can see every order, not just their own."""
    rows = db.scalars(select(ProductionWorkOrder).order_by(ProductionWorkOrder.id.desc())).all()
    return [_serialize(x) for x in rows]


@router.get("/record/{record_id}")
def get_order(record_id: int, db: Session = Depends(get_db), u=Depends(get_current_user)):
    x = db.get(ProductionWorkOrder, record_id)
    if not x:
        raise HTTPException(404, "Record not found")
    return _serialize(x)


@router.put("/record/{record_id}")
def update_order(
    record_id: int,
    p: OrderSave,
    db: Session = Depends(get_db),
    u=Depends(get_current_user),
):
    x = db.get(ProductionWorkOrder, record_id)
    if not x:
        raise HTTPException(404, "Record not found")
    if not p.data.get("date"):
        p.data["date"] = date.today().isoformat()
    _validate_order_data(p.data)

    snapshot_version(
        db, record_type="production_work_order", record_id=x.id,
        payload_json=x.payload_json, label=x.order_no, status=x.status, user=u,
    )

    x.order_no = p.order_no
    x.status = p.status
    x.payload_json = json.dumps(p.data, ensure_ascii=False, default=str)
    if p.linked_reference is not None:
        x.linked_reference = p.linked_reference.strip() or None
    db.commit(); db.refresh(x)
    return _serialize(x)


@router.delete("/record/{record_id}")
def delete_order(
    record_id: int,
    db: Session = Depends(get_db),
    u=Depends(require_roles("ADMIN", "PLANNING")),
):
    x = db.get(ProductionWorkOrder, record_id)
    if not x:
        raise HTTPException(404, "Record not found")
    db.delete(x); db.commit()
    return {"ok": True}


@router.get("/record/{record_id}/versions")
def list_order_versions(record_id: int, db: Session = Depends(get_db), u=Depends(get_current_user)):
    if not db.get(ProductionWorkOrder, record_id):
        raise HTTPException(404, "Record not found")
    rows = db.scalars(
        select(RecordVersion)
        .where(RecordVersion.record_type == "production_work_order", RecordVersion.record_id == record_id)
        .order_by(RecordVersion.id.desc())
    ).all()
    return [serialize_version(v) for v in rows]


@router.get("/record/{record_id}/versions/{version_id}")
def get_order_version(record_id: int, version_id: int, db: Session = Depends(get_db), u=Depends(get_current_user)):
    if not db.get(ProductionWorkOrder, record_id):
        raise HTTPException(404, "Record not found")
    v = db.get(RecordVersion, version_id)
    if not v or v.record_type != "production_work_order" or v.record_id != record_id:
        raise HTTPException(404, "Version not found")
    out = serialize_version(v)
    out["data"] = json.loads(v.payload_json or "{}")
    return out


@router.post("/record/{record_id}/versions/{version_id}/restore")
def restore_order_version(record_id: int, version_id: int, db: Session = Depends(get_db), u=Depends(get_current_user)):
    x = db.get(ProductionWorkOrder, record_id)
    if not x:
        raise HTTPException(404, "Record not found")
    v = db.get(RecordVersion, version_id)
    if not v or v.record_type != "production_work_order" or v.record_id != record_id:
        raise HTTPException(404, "Version not found")

    snapshot_version(
        db, record_type="production_work_order", record_id=x.id,
        payload_json=x.payload_json, label=x.order_no, status=x.status, user=u,
    )
    x.payload_json = v.payload_json
    db.commit(); db.refresh(x)
    return _serialize(x)


# ---------------------------------------------------------------- Excel export ----

_THIN = Side(style="thin", color="999999")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEAD_FONT = Font(bold=True)
_TITLE_FONT = Font(bold=True, size=14)
_WRAP = Alignment(wrap_text=True, vertical="top")


def _label_value(ws, row, label, value):
    ws[f"A{row}"] = label
    ws[f"A{row}"].font = _HEAD_FONT
    ws[f"C{row}"] = value or ""
    ws[f"C{row}"].alignment = _WRAP


def _build_workbook(order_no: str, data: dict) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "ใบสั่งผลิต"
    for col, width in zip("ABCDEFG", [4, 4, 60, 14, 10, 20, 20]):
        ws.column_dimensions[col].width = width

    ws.merge_cells("A1:G1")
    ws["A1"] = "ใบสั่งผลิต (ผลิตจริง)"
    ws["A1"].font = _TITLE_FONT
    ws.row_dimensions[1].height = 56
    _add_logo_if_present(ws, "H1")

    r = 3
    _label_value(ws, r, "เลขที่ใบสั่งผลิต", order_no); r += 1
    _label_value(ws, r, "เลขที่งาน", data.get("job_no")); r += 1
    _label_value(ws, r, "ชื่อผลิตภัณฑ์", data.get("product_name")); r += 1
    _label_value(ws, r, "ใบเสนอราคา", data.get("qp_ref")); r += 1
    _label_value(ws, r, "เลขที่สูตร", data.get("formula_ref")); r += 1
    _label_value(ws, r, "รหัสลูกค้า", data.get("customer_code")); r += 1
    _label_value(ws, r, "เลขที่ อย.", data.get("fda_no")); r += 1
    _label_value(ws, r, "LOT", data.get("lot_no")); r += 1
    _label_value(ws, r, "MFG", data.get("mfg_date")); r += 1
    _label_value(ws, r, "EXP", data.get("exp_date")); r += 1
    _label_value(ws, r, "รหัสสินค้า", data.get("product_code")); r += 1
    _label_value(ws, r, "Packing code", data.get("packing_code")); r += 1
    if data.get("production_qty"):
        _label_value(ws, r, "จำนวนผลิต", f"{data.get('production_qty')} {data.get('production_qty_unit') or ''}".strip()); r += 1
    if data.get("mg_per_unit"):
        _label_value(ws, r, "น้ำหนัก/ปริมาณต่อหน่วย", f"{data.get('mg_per_unit')} มก."); r += 1
    if data.get("bottling_qty"):
        _label_value(ws, r, "แบ่งบรรจุลง", f"{data.get('bottling_qty')} {data.get('bottling_unit') or ''}".strip()); r += 1
    if data.get("qty_per_bottle"):
        _label_value(ws, r, "ปริมาณต่อ 1 หน่วยบรรจุ", data.get("qty_per_bottle")); r += 1
    if data.get("overproduced_qty"):
        _label_value(ws, r, "จำนวนที่ผลิตเกิน", data.get("overproduced_qty")); r += 1
    _label_value(ws, r, "หมายเหตุ", data.get("notes")); r += 1
    if data.get("packing_summary"):
        _label_value(ws, r, "รายละเอียดการบรรจุ", data.get("packing_summary")); r += 1

    for group in (data.get("packing_groups") or []):
        r += 1
        ws.merge_cells(f"A{r}:G{r}")
        ws[f"A{r}"] = group.get("label") or "บรรจุแบบ"
        ws[f"A{r}"].font = _HEAD_FONT
        r += 1
        headers = ["ลำดับ", "รายละเอียด", "สเปค/สี/รหัส", "จำนวน", "หน่วย", "หมายเหตุ"]
        for i, h in enumerate(headers, start=1):
            c = ws.cell(row=r, column=i, value=h)
            c.font = _HEAD_FONT
            c.border = _BORDER
        r += 1
        for i, it in enumerate(group.get("items") or [], start=1):
            if not str(it.get("description") or "").strip():
                continue
            vals = [i, it.get("description") or "", it.get("detail") or "",
                    it.get("qty") or "", it.get("unit") or "", it.get("unit_note") or ""]
            for col, v in enumerate(vals, start=1):
                c = ws.cell(row=r, column=col, value=v)
                c.border = _BORDER
                if col == 2:
                    c.alignment = _WRAP
            r += 1

    steps = data.get("workflow_steps") or []
    if steps:
        r += 1
        ws.merge_cells(f"A{r}:G{r}")
        ws[f"A{r}"] = "ขั้นตอนการทำงาน"
        ws[f"A{r}"].font = _HEAD_FONT
        r += 1
        headers = ["ลำดับ", "รายละเอียดงานและผู้รับผิดชอบ", "", "เริ่ม", "สิ้นสุด", "ผู้ปฏิบัติงาน"]
        for i, h in enumerate(headers, start=1):
            if not h:
                continue
            c = ws.cell(row=r, column=i, value=h)
            c.font = _HEAD_FONT
            c.border = _BORDER
        r += 1
        for i, s in enumerate(steps, start=1):
            vals = [i, s.get("task") or "", s.get("responsible") or "",
                    f"{s.get('start_date') or ''} {s.get('start_time') or ''}".strip(),
                    f"{s.get('end_date') or ''} {s.get('end_time') or ''}".strip(),
                    s.get("signer") or ""]
            for col, v in enumerate(vals, start=1):
                c = ws.cell(row=r, column=col, value=v)
                c.border = _BORDER
            r += 1

    lb = data.get("labor_budget") or {}
    if any(str(v or "").strip() for v in lb.values()):
        r += 1
        ws.merge_cells(f"A{r}:G{r}")
        ws[f"A{r}"] = "งบและเวลาการทำงาน (ค่าแรง)"
        ws[f"A{r}"].font = _HEAD_FONT
        r += 1
        _label_value(ws, r, "จำนวนคน/Job", lb.get("labor_count")); r += 1
        _label_value(ws, r, "วันทำงาน/Job", lb.get("work_days")); r += 1
        _label_value(ws, r, "ค่าแรง (บาท/วัน หรือ บาท/job)", lb.get("wage_per_job")); r += 1
        _label_value(ws, r, "ค่าแรงที่ได้รับ (บาท)", lb.get("wage_received")); r += 1
        _label_value(ws, r, "ต้นทุนค่าแรง (บาท)", lb.get("labor_cost")); r += 1
        _label_value(ws, r, "กำไร/ขาดทุนค่าแรง (บาท)", lb.get("profit_loss")); r += 1
        if lb.get("notes"):
            _label_value(ws, r, "หมายเหตุ", lb.get("notes")); r += 1

    def _write_req_table(title, rows, headers, keys):
        nonlocal r
        rows = [x for x in rows if any(str(x.get(k) or "").strip() for k in keys)]
        if not rows:
            return
        r += 1
        ws.merge_cells(f"A{r}:G{r}")
        ws[f"A{r}"] = title
        ws[f"A{r}"].font = _HEAD_FONT
        r += 1
        for i, h in enumerate(headers, start=1):
            c = ws.cell(row=r, column=i, value=h)
            c.font = _HEAD_FONT
            c.border = _BORDER
        r += 1
        for i, it in enumerate(rows, start=1):
            vals = [i] + [it.get(k) or "" for k in keys]
            for col, v in enumerate(vals, start=1):
                c = ws.cell(row=r, column=col, value=v)
                c.border = _BORDER
            r += 1

    _write_req_table(
        "รายการเบิกใช้วัตถุดิบ / สารสกัด", data.get("material_requisitions") or [],
        ["ลำดับ", "LOT", "รายการที่เบิก", "จำนวนที่เบิก", "หน่วย", "หมายเหตุ"],
        ["lot_no", "item_name", "qty", "unit", "note"],
    )
    _write_req_table(
        "รายการเบิกใช้บรรจุภัณฑ์", data.get("packaging_requisitions") or [],
        ["ลำดับ", "LOT", "รายการเบิกใช้บรรจุภัณฑ์", "จำนวนที่เบิก", "หน่วย", "หมายเหตุ"],
        ["lot_no", "item_name", "qty", "unit", "note"],
    )

    coating = [x for x in (data.get("coating_formula") or [])
               if any(str(x.get(k) or "").strip() for k in ("name", "lot_no", "actual_kg"))]
    if coating:
        r += 1
        ws.merge_cells(f"A{r}:G{r}")
        ws[f"A{r}"] = "สูตรคำนวณสีเคลือบ"
        ws[f"A{r}"].font = _HEAD_FONT
        r += 1
        headers = ["ลำดับ", "ชื่อสาร", "LOT", "สูตร Test %W/W", "จริง (g)", "จริง (kg)", "+10% (g)/(kg)"]
        for i, h in enumerate(headers, start=1):
            c = ws.cell(row=r, column=i, value=h)
            c.font = _HEAD_FONT
            c.border = _BORDER
        r += 1
        for i, it in enumerate(coating, start=1):
            vals = [i, it.get("name") or "", it.get("lot_no") or "", it.get("test_pct") or "",
                    it.get("actual_g") or "", it.get("actual_kg") or "",
                    f"{it.get('need10_g') or ''} / {it.get('need10_kg') or ''}"]
            for col, v in enumerate(vals, start=1):
                c = ws.cell(row=r, column=col, value=v)
                c.border = _BORDER
                if col == 2:
                    c.alignment = _WRAP
            r += 1

    sigs = data.get("signatures") or []
    if sigs:
        r += 2
        ws.merge_cells(f"A{r}:G{r}")
        ws[f"A{r}"] = "ผู้ลงนาม"
        ws[f"A{r}"].font = _HEAD_FONT
        r += 1
        for s in sigs:
            _label_value(ws, r, s.get("role_label") or "ผู้ลงนาม",
                          f"{s.get('name') or ''} — {s.get('department') or ''} ({s.get('date') or ''})")
            r += 1

    return wb


@router.get("/record/{record_id}/excel")
def export_order_excel(record_id: int, db: Session = Depends(get_db), u=Depends(get_current_user)):
    x = db.get(ProductionWorkOrder, record_id)
    if not x:
        raise HTTPException(404, "Record not found")
    data = json.loads(x.payload_json or "{}")
    wb = _build_workbook(x.order_no, data)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    safe = re.sub(r'[^A-Za-z0-9._-]+', '_', str(x.order_no or x.id)).strip('_') or str(x.id)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="ProductionWorkOrder_{safe}.xlsx"'}
    )
